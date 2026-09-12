#!/usr/bin/env python3
"""Rewrite EULA Section 9 to match the 30-day no-questions-asked decision.

SOURCE OF THE DECISION: GatewayGuard_Decisions-RefundAndTerms-2026-08-22-1510.md
  Bill, 2026-08-22: "go with 30 day."  Published wording approved in the
  marketing plan: "If Checkup is not what you expected, write to us within
  30 days of buying and we will refund you in full. You do not have to give
  a reason."

WHY THE OLD TEXT COULD NOT STAY: it said "sales are final -- with the three
exceptions below" and required a log file to prove one of them. That is
"sales are final unless you can justify it", which is the opposite of "no
questions asked". The two could not both be published.

Assert-guarded: every paragraph replaced is matched on its existing text
first. Any mismatch aborts before a single byte is written. The input file
is never modified -- a new dated file is produced.
"""

import shutil
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "Masters" / "GatewayGuard_License-2026-08-07-0726.docx"
DST = ROOT / "Masters" / "GatewayGuard_License-2026-08-24-1210.docx"

# (index, expected text fragment) -> new text. None = delete the paragraph.
PLAN = [
    # The change log at the top DESCRIBES the old Section 9. Leaving it would
    # make the document contradict itself on page one -- the summary saying
    # "sales are final", the section itself saying the opposite.
    (10, "Section 9 now states that sales are final",
     "2. The refund policy has been rewritten to 30 days, no questions asked "
     "(2026-08-24). Section 8 of the August 4 draft was an empty placeholder. "
     "The August 7 draft filled it with a sales-are-final rule carrying three "
     "named exceptions and a log-file requirement. That has been withdrawn: "
     "Section 9 now gives a full refund on request within 30 days of purchase, "
     "with no reason required and nothing for the buyer to prove."),
    (75, "sales are final",
     "Checkup and the Guide are downloadable files. If either one is not what "
     "you expected, write to us within 30 days of buying and we will refund "
     "you in full. You do not have to give a reason."),
    (76, "We will refund your purchase in full if any of these apply",
     "You do not need to prove anything, send us a log file, or let us try to "
     "fix the problem first. If you would like to tell us what went wrong we "
     "are glad to hear it, because it is how the product improves — but it "
     "is not a condition of your refund."),
    (77, "You were charged twice", None),
    (78, "Your download never arrived", None),
    (79, "Checkup will not run on your PC", None),
    (80, "send us the log file", None),
    (81, "To ask for a refund, email",
     "To ask for a refund, email support@gatewayguard.co from the address you "
     "bought with, or use the refund link in your Gumroad receipt. Give us "
     "your order number. We aim to answer within two business days."),
    (82, "Delete every copy of the product",
     "When we refund you, your license ends and you should delete the copies "
     "you have, including any backup. We are aware we cannot check this, and "
     "we are not going to try. We are asking you to be straight with us, in "
     "the same way we are being straight with you."),
    (83, "Gumroad processes the payment",
     "If you bought through Gumroad, Gumroad handles the payment and may also "
     "issue a refund under its own policy, which can run longer than our 30 "
     "days. Nothing here takes away any right you have under the consumer law "
     "where you live."),
    (84, "DECISION NEEDED", None),
]


def main():
    if not SRC.is_file():
        raise SystemExit(f"source missing: {SRC}")

    doc = Document(str(SRC))
    paras = doc.paragraphs

    # ---- verify EVERY anchor before touching anything ----
    for idx, expect, _ in PLAN:
        if idx >= len(paras):
            raise SystemExit(f"paragraph {idx} does not exist")
        actual = paras[idx].text
        if expect not in actual:
            raise SystemExit(
                f"ANCHOR MISMATCH at paragraph {idx}\n"
                f"  expected to contain: {expect!r}\n"
                f"  actually reads     : {actual[:110]!r}\n"
                f"Nothing written. The document has changed since this "
                f"script was written -- re-read it before editing."
            )
    print(f"  all {len(PLAN)} anchors verified")

    # ---- apply, replacing text in place so run formatting survives ----
    deleted = 0
    rewritten = 0
    for idx, _expect, new in PLAN:
        p = paras[idx]
        if new is None:
            el = p._element
            el.getparent().remove(el)
            deleted += 1
            continue
        # keep the first run's formatting, drop the rest, set the text
        if p.runs:
            keep = p.runs[0]
            for r in p.runs[1:]:
                r._element.getparent().remove(r._element)
            keep.text = new
        else:
            run = p.add_run(new)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
        rewritten += 1

    print(f"  {rewritten} paragraphs rewritten, {deleted} removed")

    doc.save(str(DST))
    print(f"  written: {DST.name}")

    # ---- prove it by reading the saved file back ----
    check = Document(str(DST))
    text = "\n".join(q.text for q in check.paragraphs)
    # Scope the "must be gone" check to refund wording only. The document
    # carries four OTHER `DECISION NEEDED` blocks -- Q2(d), Q5, Q9 and Q10 --
    # which are live attorney questions on unrelated sections and must survive.
    # The first version of this check looked for the bare phrase and flagged
    # all four. A guard that fails on correct content teaches you to ignore it.
    must_be_gone = [
        "sales are final",
        "send us the log file",
        "DECISION NEEDED — refund terms",
        "You were charged twice",
    ]
    must_appear = [
        "within 30 days of buying and we will refund you in full",
        "You do not have to give a reason",
        "not a condition of your refund",
    ]
    bad = [s for s in must_be_gone if s in text]
    missing = [s for s in must_appear if s not in text]
    if bad:
        raise SystemExit(f"STILL PRESENT after save: {bad}")
    if missing:
        raise SystemExit(f"MISSING after save: {missing}")
    print("  read-back verified: old wording gone, new wording present")


if __name__ == "__main__":
    main()
